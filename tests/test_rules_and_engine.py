import pandas as pd

from mealplan.engine import Planner
from mealplan.export import export_plan
from mealplan.quantities import build_order_list, perishable_mapping
from mealplan.rules import RuleEngine


def _ok(repo, store, engine, dish, slot, day_idx=1):
    d = repo.get(dish)
    return engine.check_dish(d, store.get(d.name), slot, day_idx)


def test_hard_rules(repo, store, profile):
    e = RuleEngine(profile)
    assert not _ok(repo, store, e, "Cabbage Sabzi", "Lunch").ok          # hard exclusion
    assert not _ok(repo, store, e, "Mutton Curry", "Dinner").ok          # mutton never in-house
    assert not _ok(repo, store, e, "Bhindi Sabzi", "Lunch").ok           # bhindi excluded until style confirmed
    assert not _ok(repo, store, e, "Fish Curry", "Lunch").ok             # no seafood at lunch
    assert _ok(repo, store, e, "Fish Curry", "Dinner").ok
    assert not _ok(repo, store, e, "Prawn Garlic Butter Pasta", "Dinner").ok   # shellfish excluded for shared meals
    assert not _ok(repo, store, e, "Lemon Rice", "Dinner").ok            # no rice at night
    assert _ok(repo, store, e, "Lemon Rice", "Lunch").ok
    assert not _ok(repo, store, e, "Cucumber Salad", "Lunch").ok         # no cucumber in salads
    assert not _ok(repo, store, e, "Bombay Sandwich", "Dinner").ok       # sandwiches not at dinner
    assert _ok(repo, store, e, "Bombay Sandwich", "Breakfast").ok
    assert not _ok(repo, store, e, "Tomato Cucumber Raita", "Lunch").ok  # raita never tomato+curd
    v = _ok(repo, store, e, "Carrot Raita", "Lunch")
    assert v.ok and any("portion-only" in f for f in v.flags)          # curd -> Raksha only, flagged
    assert not _ok(repo, store, e, "Kadhi", "Lunch").ok                  # curd essential in a shared main
    assert not _ok(repo, store, e, "Idli", "Breakfast", day_idx=0).ok   # soaking not feasible on day 1
    assert _ok(repo, store, e, "Idli", "Breakfast", day_idx=1).ok


def test_rice_keyword_does_not_hit_rice_flour(repo, store, profile):
    from mealplan.rules import ingredient_matches
    akki = repo.get("Akki Roti")
    rice_flour = next(i for i in akki.ingredients if "flour" in i.name.lower())
    assert not ingredient_matches("Rice", rice_flour)
    assert ingredient_matches("Rice Flour", rice_flour)


def test_plan_is_valid_and_exports(repo, store, profile, inventory, history, llm, start_date, tmp_path):
    planner = Planner(repo, store, profile, inventory, history, llm, start_date=start_date)
    plan = planner.generate(use_llm=False)
    assert plan.violations == []
    assert len(plan.slots) == 9 and all(s.dishes for s in plan.slots)
    names = [d.name for s in plan.slots for d in s.dishes]
    assert len(names) == len(set(names))                                # no repeats
    # every dish is from the repo and legal for its slot
    for s in plan.slots:
        for d in s.dishes:
            assert d.name in repo.dishes
            assert s.slot in store.get(d.name).slots
    # rotation cap
    reused = sum(d.in_last_plan for s in plan.slots for d in s.dishes)
    assert reused <= planner.rules.rotation_allowance(len(names))
    # explainability present
    assert all(s.rationale for s in plan.slots)
    path = export_plan(plan, inventory, repo, profile, history.template, out_dir=tmp_path)
    assert path.exists()
    import openpyxl
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:3] == ["Meal Plan", "Perishable Mapping", "Order List"]
    ws = wb["Meal Plan"]
    assert ws.cell(row=4, column=2).value == "Breakfast"


def test_swap_keeps_plan_valid(repo, store, profile, inventory, history, llm, start_date):
    planner = Planner(repo, store, profile, inventory, history, llm, start_date=start_date)
    plan = planner.generate(use_llm=False)
    old = plan.get(1, "Lunch").dishes[0].name
    alts = planner.alternatives(plan, 1, "Lunch", old)
    assert alts and old not in [a.dish for a in alts]
    new = planner.swap(plan, 1, "Lunch", old, alts[0].dish, regenerate_downstream=True)
    assert alts[0].dish in new.get(1, "Lunch").names()
    assert new.violations == []


def test_repair_fixes_injected_violation(repo, store, profile, inventory, history, llm, start_date):
    planner = Planner(repo, store, profile, inventory, history, llm, start_date=start_date)
    plan = planner.generate(use_llm=False)
    # inject an illegal dish (rice at dinner) via the LLM-choice path
    from mealplan.llm import PlanChoice, SlotChoice
    slots = [SlotChoice(day=s.day, slot=s.slot, dishes=s.names(), rationale="x") for s in plan.slots]
    slots[2].dishes = ["Lemon Rice", "Aloo Gobi"]        # D1 Dinner
    bad = planner._plan_from_choice(PlanChoice(slots=slots, tradeoffs=""), source="test")
    assert planner.validate(bad)
    fixed = planner.repair(bad)
    assert fixed.violations == []
    assert "Lemon Rice" not in fixed.get(1, "Dinner").names()
    assert fixed.repairs


def test_quantity_math(repo, inventory):
    plan_dishes = [("D1 Lunch", repo.get("Aloo Gobi")), ("D1 Lunch", repo.get("Rice"))]
    mapping = pd.DataFrame(perishable_mapping(plan_dishes, inventory, 2))
    gobi = mapping[mapping.canonical == "Cauliflower"].iloc[0]
    need = sum(i.per_adult for i in repo.get("Aloo Gobi").ingredients if i.canonical == "Cauliflower") * 2
    assert abs(gobi.used - need) < 0.01 and abs(gobi.left_unused - (450 - need)) < 0.01
    orders = pd.DataFrame(build_order_list(plan_dishes, inventory, 2, set()))
    assert "Turmeric" not in set(orders["item"])      # assumed pantry
    assert set(orders["category"]) <= {"Staple", "Vegetable", "Fruit", "Dairy", "Egg", "Bread", "Chicken", "Seafood", "Mutton", "Vegetable - aromatic"}
